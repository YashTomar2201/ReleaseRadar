"""One-off convenience: convert spotcheck_blind.csv to a wrapped-text .xlsx
for comfortable manual labeling. Run again any time to regenerate from the
current CSV (e.g. after an in-progress save-back-to-csv, to re-wrap).

The CSV stays canonical -- score_spotcheck.py reads the CSV, not this file.
After labeling in Excel, convert back with csv_from_spotcheck_xlsx.py.
"""
import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

SRC = r"D:\projects\non-tech\data\labels\spotcheck_blind.csv"
DST = r"D:\projects\non-tech\data\labels\spotcheck_blind.xlsx"

df = pd.read_csv(SRC, dtype=str).fillna("")

with pd.ExcelWriter(DST, engine="openpyxl") as writer:
    df.to_excel(writer, index=False, sheet_name="spotcheck")
    ws = writer.sheets["spotcheck"]

    ws.freeze_panes = "E2"  # keep id/app/rating/review_text columns pinned while scrolling right

    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = {
        "review_id": 12, "app_key": 10, "rating": 8, "review_text": 70,
        "competitor_mentioned": 14,
    }
    default_topic_width = 15

    for idx, col_name in enumerate(df.columns, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = widths.get(col_name, default_topic_width)

    review_col_idx = list(df.columns).index("review_text") + 1
    review_width = widths["review_text"]

    for row_idx in range(2, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if col_idx == review_col_idx:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                cell.alignment = Alignment(vertical="top")

        text = str(ws.cell(row=row_idx, column=review_col_idx).value or "")
        chars_per_line = max(review_width - 2, 10)
        est_lines = max(1, -(-len(text) // chars_per_line))  # ceil div
        ws.row_dimensions[row_idx].height = min(15 * est_lines, 300)

print(f"Wrote {DST} ({len(df)} rows)")
